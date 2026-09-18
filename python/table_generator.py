# python/table_generator.py
import os
import re
import math
import openpyxl
import unicodedata


def _generate_crossref_id(text):
    """Создает безопасный ID для pandoc-crossref."""
    normalized = unicodedata.normalize('NFKD', text)
    ascii_text = normalized.encode('ascii', 'ignore').decode('ascii')
    safe_id = re.sub(r'[^a-zA-Z0-9-]', '_', ascii_text.lower())
    return re.sub(r'_+', '_', safe_id).strip('_') or 'table'


def _get_metadata(wb):
    """Читает название и тэг таблицы с листа 'Info'."""
    info_sheet = None
    for ws in wb.worksheets:
        if ws.title.lower() == 'info':
            info_sheet = ws
            break

    if not info_sheet:
        return None, None

    table_name = None
    table_tag = None

    for row in info_sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        key, value = row[0], row[1]
        if key is None:
            continue

        key_str = str(key).strip().lower()
        if key_str == 'название таблицы' and value:
            table_name = str(value).strip()
        elif key_str in ('тэг', 'tag', 'id', 'идентификатор') and value:
            table_tag = str(value).strip()

    return table_name, table_tag


def generate_table_ascii(file_abs, sheet_name='Лист1'):
    """
    Чистая функция генерации.
    Принимает абсолютный путь к Excel и имя листа.
    Возвращает строку с ASCII-таблицей и подписью.
    """
    wb = openpyxl.load_workbook(file_abs, data_only=True)
    long_title, explicit_tag = _get_metadata(wb)

    try:
        ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[sheet_name]
    except (IndexError, KeyError):
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")

    if ws.max_row is None or ws.max_row == 0 or ws.max_column is None or ws.max_column == 0:
        return ": Пустая таблица\n"

    max_row = ws.max_row
    max_col = ws.max_column
    merged_ranges = ws.merged_cells.ranges

    def get_merged_info(r, c):
        for rng in merged_ranges:
            if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                is_master = (r == rng.min_row and c == rng.min_col)
                return is_master, rng.min_row, rng.min_col, rng.max_row, rng.max_col
        return True, r, c, r, c

    # Сбор данных и расчет ширины колонок
    col_widths = [3] * max_col
    raw_data = {}
    spanned_cells = []

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            is_master, min_r, min_c, max_r, max_c = get_merged_info(r, c)
            val = ""
            if is_master:
                #val = str(ws.cell(row=r, column=c).value or '').replace('\n', ' ').strip()
                cell_value = ws.cell(row=r, column=c).value
                if cell_value is None:
                    val = ""
                else:
                    val = str(cell_value).replace('\n', ' ').strip()

            raw_data[(r, c)] = {
                'val': val, 'is_master': is_master,
                'min_r': min_r, 'min_c': min_c, 'max_r': max_r, 'max_c': max_c
            }

            if is_master:
                if min_c == max_c:
                    col_widths[c - 1] = max(col_widths[c - 1], len(val))
                else:
                    spanned_cells.append((min_c, max_c, len(val)))

    # Корректировка ширины под colspan
    spanned_cells.sort(key=lambda x: (x[1] - x[0]))
    for min_c, max_c, val_len in spanned_cells:
        current_total_width = sum(col_widths[min_c - 1:max_c]) + (max_c - min_c)
        if val_len > current_total_width:
            needed_extra = val_len - current_total_width
            num_cols = max_c - min_c + 1
            extra_per_col = math.ceil(needed_extra / num_cols)
            for c_idx in range(min_c, max_c + 1):
                col_widths[c_idx - 1] += extra_per_col

    row_heights = [1] * max_row

    # Вычисление координат
    col_starts = [0] * (max_col + 1)
    current_x = 0
    for i, w in enumerate(col_widths):
        col_starts[i] = current_x
        current_x += w + 1
    col_starts[max_col] = current_x

    row_starts = [0] * (max_row + 1)
    current_y = 0
    for i, h in enumerate(row_heights):
        row_starts[i] = current_y
        current_y += h + 1
    row_starts[max_row] = current_y

    total_width = col_starts[-1] + 1
    total_height = row_starts[-1] + 1
    canvas = [[' ' for _ in range(total_width)] for _ in range(total_height)]

    # Базовая сетка
    for r in range(max_row + 1):
        y = row_starts[r]
        for x in range(total_width - 1):
            canvas[y][x] = '-'
    for c in range(max_col + 1):
        x = col_starts[c]
        for y in range(total_height - 1):
            canvas[y][x] = '|'
    for r in range(max_row + 1):
        for c in range(max_col + 1):
            canvas[row_starts[r]][col_starts[c]] = '+'

    # Затирание внутренних границ объединений
    unique_ranges = set()
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            _, min_r, min_c, max_r, max_c = get_merged_info(r, c)
            if min_r != max_r or min_c != max_c:
                unique_ranges.add((min_r, min_c, max_r, max_c))

    for min_r, min_c, max_r, max_c in unique_ranges:
        y1, y2 = row_starts[min_r - 1], row_starts[max_r]
        x1, x2 = col_starts[min_c - 1], col_starts[max_c]

        for yi in range(y1 + 1, y2):
            for xi in range(x1 + 1, x2):
                canvas[yi][xi] = ' '
            if yi in row_starts:
                for xi in range(x1 + 1, x2):
                    if canvas[yi][xi] in '-+':
                        canvas[yi][xi] = ' '
        for xi in range(x1 + 1, x2):
            if xi in col_starts:
                for yi in range(y1 + 1, y2):
                    if canvas[yi][xi] in '|+':
                        canvas[yi][xi] = ' '

    # Вставка текста
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            info = raw_data[(r, c)]
            if info['is_master']:
                y1 = row_starts[info['min_r'] - 1]
                x1 = col_starts[info['min_c'] - 1]
                x2 = col_starts[info['max_c']]
                inner_width = (x2 - x1) - 1
                padded_val = info['val'].ljust(inner_width)[:inner_width]
                for ci, ch in enumerate(padded_val):
                    canvas[y1 + 1][x1 + 1 + ci] = ch

    # Двойная линия заголовка
    header_end_row = 2
    if header_end_row <= max_row:
        hy = row_starts[header_end_row]
        for x in range(total_width - 1):
            if canvas[hy][x] == '-':
                canvas[hy][x] = '='

    # Формирование подписи
    display_name = long_title or ws.title
    crossref_id = explicit_tag if explicit_tag else _generate_crossref_id(display_name)
    caption_line = f": {display_name} {{#tbl:{crossref_id}}}"

    grid_output = '\n'.join(''.join(row).rstrip() for row in canvas)
    return f"{grid_output}\n\n{caption_line}"