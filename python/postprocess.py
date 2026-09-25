#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Заменяет стили параграфов в DOCX и задаёт ширины столбцов таблиц.
Плюс правит word/numbering.xml: маркеры списков → тире,
абзац по левому краю, первая строка — с отступом 1 см,
таб-стоп — на 1.5 см.
"""
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import sys
import os
import re
import shutil
import zipfile
from docx import Document

# --- НАСТРОЙКИ СТИЛЕЙ ---
TARGET_STYLE_NAME = "Основной текст с отступом 31"
CAPTION_STYLE_NAME = "ДОК Таблица Текст Без Нумерации"
HEADER_STYLE_NAME = "ДОК Таблица Текст Центр"
FIRST_COL_STYLE_NAME = "ДОК Таблица Текст Центр"
OTHER_CELLS_STYLE_NAME = "ДОК Таблица Текст Центр"

# --- МАРКЕР СПИСКА ---
NEW_BULLET_MARKER = "–"

# --- ОТСТУП СПИСКА (twips, 1 см = 567) ---
LIST_LEFT_TWIPS       = 0       # весь абзац — по левому краю
LIST_FIRSTLINE_TWIPS  = 567     # первая строка (с маркером) — 1 см
LIST_TAB_TWIPS        = 850     # таб-стоп для текста после маркера — 1.5 см

# --- КАРТА ТИПОВ ТАБЛИЦ ---
TABLE_TYPE_KEYWORDS = {
    "вход для проверки": "NewTable",
    "параметр": "NewTable2",
}

# --- КАРТА ШИРИН ---
TABLE_WIDTHS_PCT = {
    ("NewTable2", 2):  [30, 70],
    ("NewTable", 6):  [16, 32, 16, 12, 12, 12],
    ("NewTable", 7):  [16, 32, 16, 9, 9, 9, 9],
    ("NewTable", 8):  [16, 33, 16, 7, 7, 7, 7, 7],
    ("NewTable", 10): [15, 23, 15, 8, 6, 6, 6, 7, 7, 7],
}


# ============================================================
#  Таблицы
# ============================================================

def set_table_layout_fixed(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')


def set_table_width_percent(table, percent=100):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(percent * 50))
    tblW.set(qn('w:type'), 'pct')


def set_table_grid(table, widths_pct):
    total = sum(widths_pct)
    if total != 100:
        raise ValueError(f"Сумма ширин должна быть 100, а не {total}")

    tbl = table._tbl
    tblPr = tbl.tblPr

    set_table_width_percent(table, 100)
    set_table_layout_fixed(table)

    old_grid = tbl.find(qn('w:tblGrid'))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement('w:tblGrid')
    for pct in widths_pct:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(50 * pct))
        grid.append(gc)
    tblPr.addnext(grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx >= len(widths_pct):
                continue
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(50 * widths_pct[idx]))
            tcW.set(qn('w:type'), 'pct')


def detect_table_type(table):
    if not table.rows:
        return None, 0
    n_cols = len(table.columns)
    header_text = table.rows[0].cells[0].text.strip().lower()
    for keyword, type_name in TABLE_TYPE_KEYWORDS.items():
        if keyword in header_text:
            return type_name, n_cols
    return None, n_cols


# ============================================================
#  Стили
# ============================================================

def replace_style(paragraph, target_style):
    if paragraph.style is None:
        return False
    if paragraph.style.name in ("Обычный", "Normal", "Основной текст", "Body Text"):
        paragraph.style = target_style
        return True
    return False


# ============================================================
#  numbering.xml: маркеры + отступы + таб
# ============================================================

def fix_lists_in_numbering(docx_path,
                           new_marker=NEW_BULLET_MARKER,
                           left_twips=LIST_LEFT_TWIPS,
                           firstline_twips=LIST_FIRSTLINE_TWIPS,
                           tab_twips=LIST_TAB_TWIPS):
    """
    Правки в word/numbering.xml:
      - у bullet-уровней (lvlText = одиночный символ) символ → new_marker,
        и убирается <w:rFonts>, чтобы маркер рисовался обычным шрифтом;
      - у ВСЕХ уровней: <w:tabs> с таб-стопом на tab_twips
        и <w:ind w:left w:firstLine>.
    """
    tmp = docx_path + ".tmp"
    shutil.move(docx_path, tmp)

    replaced_count = 0
    indent_count = 0

    LVL_TEXT_RE = re.compile(r'<w:lvlText\s+w:val="([^"]*)"\s*/>')

    def process_lvl(match):
        nonlocal replaced_count, indent_count
        block = match.group(0)

        # ── 1. bullet или numbered? ──
        mtext = LVL_TEXT_RE.search(block)
        if mtext:
            val = mtext.group(1)
            if val == "" or len(val) == 1:
                block = re.sub(
                    r'(<w:lvlText\s+w:val=")[^"]*("\s*/>)',
                    lambda mm: mm.group(1) + new_marker + mm.group(2),
                    block, count=1,
                )
                block = re.sub(r'<w:rFonts[^/]*/>', '', block)
                replaced_count += 1

        # ── 2. Отступы + таб-стоп ──
        # снять старые ind / tabs
        block = re.sub(r'<w:ind[^/]*/>', '', block)
        block = re.sub(r'<w:tabs>.*?</w:tabs>', '', block, flags=re.DOTALL)

        # порядок в pPr: ... tabs, ..., ind ...  (tabs идут ДО ind)
        new_inner = (
            f'<w:tabs><w:tab w:val="left" w:pos="{tab_twips}"/></w:tabs>'
            f'<w:ind w:left="{left_twips}" w:firstLine="{firstline_twips}"/>'
        )

        ppr_match = re.search(r'<w:pPr[^>]*>', block)
        if ppr_match:
            at = ppr_match.end()
            block = block[:at] + new_inner + block[at:]
        else:
            block = re.sub(r'</w:lvl>',
                           f'<w:pPr>{new_inner}</w:pPr></w:lvl>',
                           block)
        indent_count += 1
        return block

    try:
        with zipfile.ZipFile(tmp, "r") as zin, \
             zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)

                if item.filename == "word/numbering.xml":
                    text = data.decode("utf-8")
                    text = re.sub(
                        r'<w:lvl\b.*?</w:lvl>',
                        process_lvl,
                        text,
                        flags=re.DOTALL,
                    )
                    data = text.encode("utf-8")

                zout.writestr(item, data)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass

    print(f"  Bullet markers replaced: {replaced_count} → '{new_marker}'")
    print(f"  List indents set: {indent_count} level(s), "
          f"left={left_twips}, firstLine={firstline_twips}, "
          f"tab={tab_twips} twips")
    return replaced_count, indent_count


# ============================================================
#  main
# ============================================================

def main():
    if len(sys.argv) < 2:
        print("Usage: python postprocess.py <docx_path>", file=sys.stderr)
        sys.exit(1)

    docx_path = sys.argv[1]

    if not os.path.isfile(docx_path):
        print(f"File not found: {docx_path}", file=sys.stderr)
        sys.exit(1)

    if os.environ.get("TEST") == "1":
        print(f"[TEST MODE] Skipping post-processing for {docx_path}")
        return

    print(f"Post-processing: {docx_path}")
    doc = Document(docx_path)

    available_styles = {s.name for s in doc.styles}
    required_styles = [
        TARGET_STYLE_NAME, CAPTION_STYLE_NAME,
        HEADER_STYLE_NAME, FIRST_COL_STYLE_NAME, OTHER_CELLS_STYLE_NAME,
    ]
    missing_styles = [s for s in required_styles if s not in available_styles]
    if missing_styles:
        print(f"Error: Missing styles in document: {missing_styles}", file=sys.stderr)
        print(f"Available styles: {sorted(available_styles)}", file=sys.stderr)
        sys.exit(2)

    style_body        = doc.styles[TARGET_STYLE_NAME]
    style_caption     = doc.styles[CAPTION_STYLE_NAME]
    style_header      = doc.styles[HEADER_STYLE_NAME]
    style_first_col   = doc.styles[FIRST_COL_STYLE_NAME]
    style_other_cells = doc.styles[OTHER_CELLS_STYLE_NAME]

    replaced = 0

    # 1. Основной текст
    for p in doc.paragraphs:
        if "Таблица" in p.text:
            if replace_style(p, style_caption):
                replaced += 1
        else:
            if replace_style(p, style_body):
                replaced += 1

    # 2. Таблицы
    for table in doc.tables:
        type_name, n_cols = detect_table_type(table)

        if type_name is not None:
            widths = TABLE_WIDTHS_PCT.get((type_name, n_cols))
            if widths is not None:
                set_table_grid(table, widths)
                print(f"  {type_name} ({n_cols} cols): widths = {widths}%")
            else:
                print(f"  [warn] Нет ширин для {type_name} с {n_cols} столбцами — пропуск",
                      file=sys.stderr)
                set_table_width_percent(table, 100)
        else:
            set_table_width_percent(table, 100)

        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                current_target_style = style_other_cells
                if row_idx == 0:
                    current_target_style = style_header
                elif cell_idx == 0:
                    current_target_style = style_first_col

                for p in cell.paragraphs:
                    if replace_style(p, current_target_style):
                        replaced += 1

    print(f"Replaced paragraphs: {replaced}")
    doc.save(docx_path)

    # 3. numbering.xml: маркеры + отступы + таб
    fix_lists_in_numbering(docx_path)

    print("Saved:", docx_path)
    print("Done.")


if __name__ == "__main__":
    main()