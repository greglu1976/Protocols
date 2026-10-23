#!/usr/bin/env python
# -*- coding: utf-8 -*-
import json
import os
import re
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from table_generator import generate_table_ascii


REPEAT_BY_FILE = {
    "unit.xlsx": ("code_order", "code_order2", "code_order3"),
    "hmi.xlsx": ("code_hmi", "code_hmi2", "code_hmi3"),
    "analog.xlsx": ("code_order", "code_order2", "code_order3"),
}

GATE_BY_FILE = {
    "prv.xlsx": "code_prmprd",
}

DATA_SHEET = "Лист1"
PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
DISCRET_JSON_DIR = Path(__file__).resolve().parent.parent / "discret"

DISCRET_HEADERS = (
    "Наименование задействованного входа",
    "U ср / U в, В",
    "Наименование задействованного входа",
    "U ср / U в, В",
)


def _nonempty(v) -> bool:
    if v is None: return False
    if isinstance(v, str) and not v.strip(): return False
    return True


def _substitute_in_xlsx(template_path: Path, ctx: dict):
    wb = openpyxl.load_workbook(template_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str): continue
                def repl(m):
                    val = ctx.get(m.group(1))
                    return "" if val is None else str(val)
                new = PLACEHOLDER_RE.sub(repl, v)
                if new != v: cell.value = new
    return wb


def _pick_data_sheet(wb):
    if DATA_SHEET in wb.sheetnames: return wb[DATA_SHEET]
    for s in wb.worksheets:
        if s.title.lower() != "info": return s
    return wb.worksheets[0]


def _save_and_render(wb, ws) -> str:
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(tmp_fd)
    try:
        wb.save(tmp_path)
        return generate_table_ascii(tmp_path, sheet_name=ws.title)
    finally:
        try: os.remove(tmp_path)
        except OSError: pass


# --- Аналоговые каналы ---
_CYR_TO_LAT = {
    "\u0410": "A", "\u0412": "B", "\u0421": "C", "\u0415": "E",
    "\u041d": "H", "\u041a": "K", "\u041c": "M", "\u041e": "O",
    "\u0420": "P", "\u0422": "T", "\u0425": "X",
}

def _char_value(ch: str) -> int:
    if ch.isdigit(): return int(ch)
    up = ch.upper()
    up = _CYR_TO_LAT.get(up, up)
    if "A" <= up <= "Z": return 10 + (ord(up) - ord("A"))
    return 0

def calculate_analog_rows(code: str) -> int:
    if not code: return 0
    total = 0
    for part in re.split(r"[-\u2010\u2011\u2012\u2013\u2014\u2015]", code):
        m = re.match(r"^[Mm\u041c\u043c]1([^.]*)", part)
        if not m: continue
        total += sum(_char_value(ch) for ch in m.group(1))
    return total


# --- Дискретные входы (Grid Table с символом \) ---
def _terminals_to_str(value) -> str:
    if isinstance(value, (list, tuple)):
        s = ", ".join(str(x) for x in value)
    else:
        s = "" if value is None else str(value)
    return s.replace("\u0425", "X").replace("\u0445", "x")


def _make_grid_table(header, rows):
    """Вспомогательная функция для рисования сетчатой таблицы."""
    widths = [len(h) for h in header]
    for row in rows:
        for i, cell in enumerate(row):
            for line in cell.split('\n'):
                widths[i] = max(widths[i], len(line))
    
    def border(char='-'):
        return '+' + '+'.join(char * (w + 2) for w in widths) + '+'
    
    def format_row(cells):
        max_lines = max(len(c.split('\n')) for c in cells)
        out_lines = []
        for ln in range(max_lines):
            parts = []
            for i, cell in enumerate(cells):
                lines = cell.split('\n')
                val = lines[ln] if ln < len(lines) else ""
                parts.append(f" {val.ljust(widths[i])} ")
            out_lines.append('|' + '|'.join(parts) + '|')
        return '\n'.join(out_lines)

    md = [border('-'), format_row(header), border('=')]
    for row in rows:
        md.append(format_row(row))
        md.append(border('-'))
    return '\n'.join(md)


def render_discret_from_json(vars_data: dict) -> str:
    cabinet_id = str(vars_data.get("id", ""))
    json_path = None

    if DISCRET_JSON_DIR.is_dir():
        for p in DISCRET_JSON_DIR.glob("*.json"):
            if cabinet_id in p.name:
                json_path = p
                break

    if not json_path or not json_path.is_file():
        print("[general_updater] WARNING: JSON для {} не найден".format(cabinet_id), file=sys.stderr)
        return ""

    try:
        with open(json_path, "r", encoding="utf-8") as f:
            raw_data = json.load(f).get("data", {})
    except Exception as e:
        print("[general_updater] ERROR reading {}: {}".format(json_path, e), file=sys.stderr)
        return ""

    tables = []
    # Проходим по каждому устройству (A1, A2, A3)
    for device_name, device_data in raw_data.items():
        if not isinstance(device_data, dict) or not device_data:
            continue
            
        items = list(device_data.items())
        mid_idx = (len(items) + 1) // 2
        col1_items = items[:mid_idx]
        col2_items = items[mid_idx:]
        
        rows = []
        for i in range(max(len(col1_items), len(col2_items))):
            c1, c2 = "", ""
            if i < len(col1_items):
                k1, v1 = col1_items[i]
                # \ (обратный слэш) в конце строки — команда Pandoc сделать LineBreak
                c1 = f"{k1} \\\n{_terminals_to_str(v1)}"
            if i < len(col2_items):
                k2, v2 = col2_items[i]
                c2 = f"{k2} \\\n{_terminals_to_str(v2)}"
            rows.append([c1, "", c2, ""])

        table_str = _make_grid_table(DISCRET_HEADERS, rows)
        
        # Делаем уникальный ID для каждой таблицы (например, tbl:discret_A1)
        crossref_id = f"discret_{device_name}"
        caption = f"\n\n: Пороги срабатывания дискретных входов ({device_name}) {{#tbl:{crossref_id}}}\n\n"
        
        tables.append(f"\n\n{table_str}{caption}\n\n")

    return "\n".join(tables)


# --- Размножение unit/hmi/analog ---
def _expand_repeat(fname: str, vars_data: dict, templates_dir: Path) -> str:
    keys = REPEAT_BY_FILE[fname]
    tables = []
    for idx, key in enumerate(keys, start=1):
        code = vars_data.get(key)
        if not _nonempty(code): continue

        ctx = dict(vars_data)
        ctx["code_order"] = code
        ctx["code_hmi"] = vars_data.get("code_hmi{}".format(idx) if idx > 1 else "code_hmi")
        ctx["label"] = "A{}".format(idx)
        ctx["N"] = str(idx)

        n_extra = 0
        if fname == "analog.xlsx":
            n_extra = calculate_analog_rows(code)

        wb = _substitute_in_xlsx(templates_dir / fname, ctx)
        ws = _pick_data_sheet(wb)

        if n_extra:
            start = ws.max_row + 1
            for r in range(start, start + n_extra):
                ws.cell(row=r, column=1, value=" ")

        try:
            md = _save_and_render(wb, ws)
            tables.append(md)
        except Exception as e:
            print("[general_updater] ERROR rendering {} (N={}): {}".format(fname, idx, e), file=sys.stderr)

    if not tables: return ""
    return "\n\n".join(tables) + "\n"


# --- Блоки INSERT_TABLE ---
BLOCK_RE = re.compile(
    r"<!--\s*INSERT_TABLE:file=(?P<file>[^\s,>]+)(?:[^>]*)?\s*-->"
    r"(?P<body>.*?)"
    r"<!--\s*END_INSERT_TABLE\s*-->",
    re.DOTALL,
)

def process_general_md(template_text: str, vars_data: dict, templates_dir: Path) -> str:
    def repl(m):
        fname = m.group("file")
        if fname in REPEAT_BY_FILE:
            return _expand_repeat(fname, vars_data, templates_dir)
        if fname == "discret.xlsx":
            return render_discret_from_json(vars_data)
            
        gate = GATE_BY_FILE.get(fname)
        if gate is not None and not _nonempty(vars_data.get(gate)):
            return ""

        xlsx_path = templates_dir / fname
        if not xlsx_path.is_file(): return m.group(0)

        wb = _substitute_in_xlsx(xlsx_path, dict(vars_data))
        ws = _pick_data_sheet(wb)
        try:
            md = _save_and_render(wb, ws)
        except Exception as e:
            print("[general_updater] ERROR rendering {}: {}".format(fname, e), file=sys.stderr)
            return m.group(0)
        return md + "\n"

    return BLOCK_RE.sub(repl, template_text)


# --- main ---
def main():
    if len(sys.argv) != 5:
        print("Usage: general_updater.py <vars.json> <template_general.md> <out_general.md> <templates_dir>", file=sys.stderr)
        return 2

    vars_json = Path(sys.argv[1])
    template_md = Path(sys.argv[2])
    out_md = Path(sys.argv[3])
    templates_dir = Path(sys.argv[4])

    out_md.parent.mkdir(parents=True, exist_ok=True)

    vars_data = json.loads(vars_json.read_text(encoding="utf-8"))
    template_text = template_md.read_text(encoding="utf-8")

    result = process_general_md(template_text, vars_data, templates_dir)
    out_md.write_text(result, encoding="utf-8", newline="\n")

    print("[general_updater] wrote {}".format(out_md))
    return 0


if __name__ == "__main__":
    sys.exit(main())