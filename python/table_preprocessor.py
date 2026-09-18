#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Препроцессор таблиц для протоколов.
Поддерживает парные теги <!-- INSERT_TABLE --> ... <!-- END_INSERT_TABLE -->
для безопасного обновления таблиц без дублирования.
"""
import sys
import os
import re
import math
import openpyxl
from pathlib import Path
import unicodedata

# Регулярка для парсинга параметров внутри открывающего тега
TAG_PARAMS_RE = re.compile(
    r'(?:file=(?P<file>[^,\s]+),?\s*)?'           
    r'(?:,\s*sheet(?:_name)?=(?P<sheet>[^,]+?))?'  
)

OPEN_TAG = "<!-- INSERT_TABLE:"
CLOSE_TAG = "<!-- END_INSERT_TABLE -->"


def generate_crossref_id(text):
    """Создает безопасный ID для pandoc-crossref."""
    normalized = unicodedata.normalize('NFKD', text)
    ascii_text = normalized.encode('ascii', 'ignore').decode('ascii')
    safe_id = re.sub(r'[^a-zA-Z0-9-]', '_', ascii_text.lower())
    return re.sub(r'_+', '_', safe_id).strip('_') or 'table'


def get_table_metadata(wb):
    """Читает название и тэг из листа 'Info'."""
    info_sheet = None
    for ws in wb.worksheets:
        if ws.title.lower() == 'info':
            info_sheet = ws
            break
            
    if not info_sheet:
        return None, None
        
    table_name = None
    table_tag = None
    
    for row in info_sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        key, value = row[0], row[1]
        if key is None: continue
            
        key_str = str(key).strip().lower()
        if key_str == 'название таблицы' and value:
            table_name = str(value).strip()
        elif key_str in ('тэг', 'tag', 'id', 'идентификатор') and value:
            table_tag = str(value).strip()
            
    return table_name, table_tag


def excel_to_perfect_grid(file_path, sheet_name='Лист1'):
    """Генерирует ASCII-сетку + подпись."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    long_title, explicit_tag = get_table_metadata(wb)
    
    try:
        ws = wb[sheet_name] if isinstance(sheet_name, str) else wb.worksheets[sheet_name]
    except (IndexError, KeyError):
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
        
    if ws.max_row is None or ws.max_row == 0 or ws.max_column is None or ws.max_column == 0:
        return ": Пустая таблица\n"

    max_row = ws.max_row
    max_col = ws.max_column
    merged_ranges = ws.merged_cells.ranges

    def get_merged_info(r, c):
        for rng in merged_ranges:
            if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                return (r == rng.min_row and c == rng.min_col), \
                       rng.min_row, rng.min_col, rng.max_row, rng.max_col
        return True, r, c, r, c

    col_widths = [3] * max_col
    raw_data = {}
    spanned_cells = []

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            is_master, min_r, min_c, max_r, max_c = get_merged_info(r, c)
            val = ""
            if is_master:
                val = str(ws.cell(row=r, column=c).value or '').replace('\n', ' ').strip()

            raw_data[(r, c)] = {
                'val': val, 'is_master': is_master,
                'min_r': min_r, 'min_c': min_c, 'max_r': max_r, 'max_c': max_c
            }

            if is_master:
                if min_c == max_c:
                    col_widths[c - 1] = max(col_widths[c - 1], len(val))
                else:
                    spanned_cells.append((min_c, max_c, len(val)))

    spanned_cells.sort(key=lambda x: (x[1] - x[0]))
    for min_c, max_c, val_len in spanned_cells:
        current_total_width = sum(col_widths[min_c - 1:max_c]) + (max_c - min_c)
        if val_len > current_total_width:
            needed_extra = val_len - current_total_width
            num_cols = max_c - min_c + 1
            extra_per_col = math.ceil(needed_extra / num_cols)
            for c_idx in range(min_c, max_c + 1):
                col_widths[c_idx - 1] += extra_per_col

    row_heights = [1] * max_row
    col_starts = [0] * (max_col + 1)
    current_x = 0
    for i, w in enumerate(col_widths):
        col_starts[i] = current_x
        current_x += w + 1
    col_starts[max_col] = current_x

    row_starts = [0] * (max_row + 1)
    current_y = 0
    for i, h in enumerate(row_heights):
        row_starts[i] = current_y
        current_y += h + 1
    row_starts[max_row] = current_y

    total_width = col_starts[-1] + 1
    total_height = row_starts[-1] + 1
    canvas = [[' ' for _ in range(total_width)] for _ in range(total_height)]

    for r in range(max_row + 1):
        y = row_starts[r]
        for x in range(total_width - 1): canvas[y][x] = '-'
    for c in range(max_col + 1):
        x = col_starts[c]
        for y in range(total_height - 1): canvas[y][x] = '|'
    for r in range(max_row + 1):
        for c in range(max_col + 1): canvas[row_starts[r]][col_starts[c]] = '+'

    unique_ranges = set()
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            _, min_r, min_c, max_r, max_c = get_merged_info(r, c)
            if min_r != max_r or min_c != max_c:
                unique_ranges.add((min_r, min_c, max_r, max_c))

    for min_r, min_c, max_r, max_c in unique_ranges:
        y1, y2 = row_starts[min_r - 1], row_starts[max_r]
        x1, x2 = col_starts[min_c - 1], col_starts[max_c]
        for yi in range(y1 + 1, y2):
            for xi in range(x1 + 1, x2): canvas[yi][xi] = ' '
            if yi in row_starts:
                for xi in range(x1 + 1, x2):
                    if canvas[yi][xi] in '-+': canvas[yi][xi] = ' '
        for xi in range(x1 + 1, x2):
            if xi in col_starts:
                for yi in range(y1 + 1, y2):
                    if canvas[yi][xi] in '|+': canvas[yi][xi] = ' '

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            info = raw_data[(r, c)]
            if info['is_master']:
                y1 = row_starts[info['min_r'] - 1]
                x1 = col_starts[info['min_c'] - 1]
                x2 = col_starts[info['max_c']]
                inner_width = (x2 - x1) - 1
                padded_val = info['val'].ljust(inner_width)[:inner_width]
                for ci, ch in enumerate(padded_val):
                    canvas[y1 + 1][x1 + 1 + ci] = ch

    header_end_row = 2
    if header_end_row <= max_row:
        hy = row_starts[header_end_row]
        for x in range(total_width - 1):
            if canvas[hy][x] == '-': canvas[hy][x] = '='

    display_name = long_title or ws.title
    crossref_id = explicit_tag if explicit_tag else generate_crossref_id(display_name)
    caption_line = f": {display_name} {{#tbl:{crossref_id}}}"

    grid_output = '\n'.join(''.join(row).rstrip() for row in canvas)
    return f"{grid_output}\n\n{caption_line}"


def process_file(md_path):
    """Обрабатывает файл, находя ПАРНЫЕ теги и обновляя блоки."""
    content = Path(md_path).read_text(encoding='utf-8')
    md_file_dir = os.path.dirname(os.path.abspath(md_path))
    
    result_parts = []
    pos = 0
    updated_count = 0
    
    while pos < len(content):
        # Ищем следующий открывающий тег
        open_idx = content.find(OPEN_TAG, pos)
        
        if open_idx == -1:
            # Тегов больше нет, добавляем остаток текста
            result_parts.append(content[pos:])
            break
            
        # Добавляем текст ДО открывающего тега
        result_parts.append(content[pos:open_idx])
        
        # Находим конец открывающего тега
        open_end_idx = content.find('-->', open_idx)
        if open_end_idx == -1:
            # Невалидный тег, пропускаем
            result_parts.append(content[open_idx:])
            break
            
        # Парсим параметры
        tag_content = content[open_idx + len(OPEN_TAG):open_end_idx]
        params_match = TAG_PARAMS_RE.match(tag_content.strip())
        params = params_match.groupdict() if params_match else {}
        
        file_rel = params.get('file')
        if not file_rel:
            print("[PREPROCESS] ERROR: Missing 'file=' in tag", file=sys.stderr)
            block = content[open_idx:open_end_idx + 3]
            result_parts.append(block)
            pos = open_end_idx + 3
            continue
            
        sheet_name = params.get('sheet') or 'Лист1'
        file_abs = os.path.normpath(os.path.join(md_file_dir, file_rel))
        
        # Ищем ЗАКРЫВАЮЩИЙ тег
        close_idx = content.find(CLOSE_TAG, open_end_idx)
        if close_idx == -1:
            print(f"[PREPROCESS] WARNING: No closing tag for table at pos {open_idx}", file=sys.stderr)
            # Если закрывающего тега нет, просто оставляем открывающий как есть
            result_parts.append(content[open_idx:open_end_idx + 3])
            pos = open_end_idx + 3
            continue
            
        # Генерируем новую таблицу
        if os.path.isfile(file_abs):
            try:
                new_table = excel_to_perfect_grid(file_abs, sheet_name=sheet_name)
                table_block = f"\n{new_table}\n"
                updated_count += 1
                print(f"[PREPROCESS] Updated table from '{file_rel}'")
            except Exception as e:
                import traceback
                print(f"[PREPROCESS] ERROR processing '{file_rel}':", file=sys.stderr)
                traceback.print_exc()
                table_block = f"\n[ERROR: Failed to process '{file_rel}']\n"
        else:
            print(f"[PREPROCESS] ERROR: File not found: {file_abs}", file=sys.stderr)
            table_block = f"\n[ERROR: File '{file_rel}' not found]\n"
        
        # Собираем блок: ОТКРЫВАЮЩИЙ ТЕГ + ТАБЛИЦА + ЗАКРЫВАЮЩИЙ ТЕГ
        full_open_tag = content[open_idx:open_end_idx + 3]
        result_parts.append(full_open_tag)
        result_parts.append(table_block)
        result_parts.append(CLOSE_TAG)
        
        # Перемещаем позицию ПОСЛЕ закрывающего тега
        pos = close_idx + len(CLOSE_TAG)
    
    new_content = ''.join(result_parts)
    
    if updated_count > 0:
        Path(md_path).write_text(new_content, encoding='utf-8')
        
    return new_content


def main():
    if len(sys.argv) < 2:
        print("Usage: python table_preprocessor.py file1.md [file2.md ...]", file=sys.stderr)
        sys.exit(1)

    processed_count = 0
    for md_file in sys.argv[1:]:
        if not os.path.isfile(md_file):
            print(f"[PREPROCESS] Warning: File not found: {md_file}", file=sys.stderr)
            continue

        process_file(md_file)
        processed_count += 1

    print(f"[PREPROCESS] Done. Processed {processed_count} file(s).")


if __name__ == "__main__":
    main()